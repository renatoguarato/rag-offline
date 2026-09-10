from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.application.ports import ProcessedDocument


class DocumentProcessingAdapter:
    CHUNK_SIZE = 1000
    CHUNK_OVERLAP = 200

    def process(self, content_type: str, content: bytes) -> ProcessedDocument:
        content_type = content_type.lower()
        if "pdf" in content_type:
            text = "\n".join(
                page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages
            )
        elif "docx" in content_type or "word" in content_type:
            text = "\n".join(
                paragraph.text for paragraph in DocxDocument(BytesIO(content)).paragraphs
            )
        elif "text" in content_type or "txt" in content_type:
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
        else:
            raise ValueError(f"Unsupported content type: {content_type}")
        if not text.strip():
            raise ValueError("Extracted text is empty")
        return ProcessedDocument(self._chunk(text))

    def _chunk(self, text: str) -> list[str]:
        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = start + self.CHUNK_SIZE
            if end >= len(text):
                chunks.append(text[start:].strip())
                break
            chunk = text[start:end]
            split_point = max(chunk.rfind("."), chunk.rfind("\n"))
            if split_point > self.CHUNK_SIZE // 2:
                end = start + split_point + 1
                chunk = text[start:end]
            chunks.append(chunk.strip())
            start = end - self.CHUNK_OVERLAP
        return chunks
